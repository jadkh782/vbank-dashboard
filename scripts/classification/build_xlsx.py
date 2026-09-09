import json, collections, datetime as dt, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from family import family_key

SRC = sys.argv[1]
OUT = sys.argv[2]
d = json.load(open(SRC, encoding='utf-8'))
E = d['errors']

CATS = ['Exelentic', 'V-Bank IT', 'Neustartfähiger Vorgang', 'Nicht als Fehler anzeigen']
CUR_MAP = {'Exelentic': 'Exelentic', 'V-Bank IT': 'V-Bank IT',
           'Neustartfähiger Vorgang': 'Neustartfähiger Vorgang',
           'Kein Fehler (Business)': 'Nicht als Fehler anzeigen'}

def cls(source):
    if source.startswith('Business'): return 'biz'
    if source.startswith('Job'): return 'job'
    return 'sys'

TYP = {'biz': 'Business Exception (Warteschlange)', 'job': 'Job abgebrochen (Prozess)',
       'sys': 'System Exception (Warteschlange)'}

# ── build families ──────────────────────────────────────────────────────────
fam = {}
for e in E:
    c = cls(e['source'])
    k = (c, family_key(e['message'], e['source']))
    f = fam.get(k)
    if not f:
        f = fam[k] = {'cls': c, 'key': k[1], 'count': 0, 'variants': [], 'first': e['firstSeen'],
                      'last': e['lastSeen'], 'procs': set(), 'folders': set(), 'cur': collections.Counter()}
    f['count'] += e['count']
    f['variants'].append(e)
    f['first'] = min(f['first'], e['firstSeen'])
    f['last'] = max(f['last'], e['lastSeen'])
    f['procs'].update(e['processes'])
    f['folders'].update(e['folders'])
    f['cur'][CUR_MAP[e['currentDashboard']]] += e['count']

fams = sorted(fam.values(), key=lambda f: -f['count'])
_n = {'main': 0, 'biz': 0}
for f in fams:
    if f['cls'] == 'biz':
        _n['biz'] += 1; f['nr'] = f"B{_n['biz']}"
    else:
        _n['main'] += 1; f['nr'] = _n['main']
    f['variants'].sort(key=lambda v: -v['count'])
    top = f['variants'][0]
    f['message'] = top['message']
    f['sample'] = top['sample']
    cur = f['cur'].most_common()
    f['current'] = cur[0][0] if len(cur) == 1 else f'{cur[0][0]} (gemischt: ' + ', '.join(f'{k} {v}' for k, v in cur) + ')'

main = [f for f in fams if f['cls'] != 'biz']
biz = [f for f in fams if f['cls'] == 'biz']
fam_of = {}
for f in fams:
    for v in f['variants']:
        fam_of[id(v)] = f['nr']

# ── styles ──────────────────────────────────────────────────────────────────
FONT = 'Arial'
base = Font(name=FONT, size=10)
bold = Font(name=FONT, size=10, bold=True)
hdr_font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='1F3864')
input_fill = PatternFill('solid', fgColor='FFFF00')
input_hdr_fill = PatternFill('solid', fgColor='C65911')
grey = Font(name=FONT, size=9, color='666666')
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
top = Alignment(vertical='top')

def iso(s):
    if not s: return None
    return dt.datetime.fromisoformat(s.replace('Z', '+00:00')).replace(tzinfo=None)

def header(ws, cols, input_cols=(), freeze='C2'):
    for c, (name, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = hdr_font
        cell.fill = input_hdr_fill if name in input_cols else hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = freeze

def style_row(ws, r, ncols, input_cols_idx=()):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = base
        cell.border = border
        cell.alignment = wrap if c not in input_cols_idx else Alignment(vertical='top')
        if c in input_cols_idx:
            cell.fill = input_fill

wb = Workbook()

# ── Sheet: Anleitung ────────────────────────────────────────────────────────
ws = wb.active
ws.title = 'Anleitung'
ws.column_dimensions['A'].width = 34
ws.column_dimensions['B'].width = 110
lines = [
    ('Fehlerklassifizierung V-Bank Dashboard', None),
    ('Datenquelle', f"Orchestrator {d['orchestrator']}, alle {d['totals']['folders']} Ordner, gesamte Historie (kein Zeitfenster). Export: {d['exportedAt'][:16].replace('T', ' ')} UTC"),
    ('Umfang', f"{d['totals']['jobs']:,} Jobs (davon {d['totals']['faultedJobs']:,} abgebrochen), {d['totals']['failedQueueItems']:,} fehlgeschlagene/erneut versuchte/abgebrochene Queue-Items, {d['totals']['processes']} Prozesse, {d['totals']['queues']} Warteschlangen.".replace(',', '.')),
    ('', None),
    ('Was ist zu tun', 'Gelbe Zellen ausfüllen. Alle anderen Spalten sind Daten aus dem Orchestrator und müssen nicht geändert werden.'),
    ('Blatt "Fehler"', f'{len(main)} Fehlerfamilien (System Exceptions auf Queue-Items und abgebrochene Jobs), nach Häufigkeit sortiert. Pro Zeile in Spalte "Neue Einstufung" eine Kategorie wählen. Spalte "Anteil kumuliert" zeigt, wie viel Volumen die Zeilen bis hierher abdecken – die ersten ~100 Zeilen decken den Großteil ab.'),
    ('Blatt "Business Exceptions"', f'{len(biz)} Familien von Business Exceptions. Laut Fachregel sind das keine Störungen (korrekt erkannte Aussteuerung), deshalb ist "Nicht als Fehler anzeigen" bereits vorbelegt. Nur ändern, wo das nicht stimmt.'),
    ('Blatt "Prozesse & Warteschlangen"', 'Alle Prozesse (Releases) und Warteschlangen mit Ordner und Aktivität. In Spalte "Einbeziehen" Ja/Nein wählen: Nein = wird im Dashboard komplett ausgeblendet.'),
    ('Blatt "Varianten (Detail)"', 'Alle einzelnen Fehlermeldungen (normalisiert) mit Zuordnung zur Familien-Nr. Nur zum Nachschlagen; wird bei der Umsetzung im Code verwendet.'),
    ('', None),
    ('Kategorien', None),
    ('Exelentic', 'Defekt in der Automatisierung selbst (Logik, Selektoren, Datenverarbeitung). Exelentic muss handeln.'),
    ('V-Bank IT', 'Infrastruktur: Server, Netzwerk, Zugänge, Drittsysteme (Avaloq, Citrix, Laufwerke, Mail). V-Bank IT muss handeln.'),
    ('Neustartfähiger Vorgang', 'Vorübergehender Fehler auf einem einzelnen Vorgang; wird erneut gestartet und läuft dann normal durch. Niemand muss handeln, wird aber als offener Punkt gezählt.'),
    ('Nicht als Fehler anzeigen', 'Erscheint nicht als Störung im Dashboard (z. B. Business Exception, Testlauf, bewusst gestoppter Job).'),
    ('', None),
    ('Spalte "Aktuelle Einstufung"', 'So stuft das Dashboard die Meldung heute ein (Schlüsselwortregel: ' + ', '.join(['server', 'timeout', 'connection', 'network', 'login', 'unavailable', 'remote', 'disconnected', '502', '503', 'crashed']) + ' → V-Bank IT). "gemischt" = die Varianten der Familie werden heute unterschiedlich eingestuft.'),
    ('Spalte "Muster"', 'Technischer Gruppierungsschlüssel: Namen, Nummern, Dateinamen, Selektoren und Stacktraces sind durch Platzhalter ersetzt, damit gleiche Fehler mit unterschiedlichen Werten eine Zeile bilden.'),
    ('Spalte "Varianten"', 'Anzahl unterschiedlicher normalisierter Meldungen in der Familie (siehe Blatt "Varianten (Detail)").'),
]
for r, (a, b) in enumerate(lines, 1):
    ca = ws.cell(row=r, column=1, value=a)
    ca.font = Font(name=FONT, size=14, bold=True) if r == 1 else bold
    ca.alignment = top
    if b is not None:
        cb = ws.cell(row=r, column=2, value=b)
        cb.font = base
        cb.alignment = wrap
    if a in CATS:
        ca.fill = input_fill
ws.cell(row=len(lines) + 2, column=1, value='Zusammenfassung').font = bold
sr = len(lines) + 3
for lbl, formula in [
    ('Fehlerfamilien (Blatt Fehler)', '=COUNTA(Fehler!A:A)-1'),
    ('davon bereits eingestuft', '=COUNTA(Fehler!N:N)-1'),
    ('Business-Exception-Familien', "=COUNTA('Business Exceptions'!A:A)-1"),
    ('Prozesse + Warteschlangen', "=COUNTA('Prozesse & Warteschlangen'!B:B)-1"),
    ('davon mit Entscheidung Einbeziehen', "=COUNTA('Prozesse & Warteschlangen'!L:L)-1"),
]:
    ws.cell(row=sr, column=1, value=lbl).font = base
    ws.cell(row=sr, column=2, value=formula).font = base
    sr += 1

# ── Sheet: Fehler / Business Exceptions ─────────────────────────────────────
ERR_COLS = [
    ('Nr', 6), ('Typ', 20), ('Meldung (häufigste Variante)', 70), ('Muster', 45), ('Beispiel (Original)', 60),
    ('Anzahl', 9), ('Anteil kumuliert', 10), ('Varianten', 9), ('Erstes Auftreten', 14), ('Letztes Auftreten', 14),
    ('Prozesse / Warteschlangen', 40), ('Ordner', 30), ('Aktuelle Einstufung (Dashboard heute)', 24),
    ('Neue Einstufung', 26), ('Kommentar', 40),
]
dv_cat = DataValidation(type='list', formula1='"' + ','.join(CATS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Kategorie', error='Bitte eine der vier Kategorien wählen.')

def error_sheet(title, rows, prefill=None):
    ws = wb.create_sheet(title)
    header(ws, ERR_COLS, input_cols=('Neue Einstufung', 'Kommentar'))
    ws.add_data_validation(dv_cat)
    n = len(rows)
    for i, f in enumerate(rows):
        r = i + 2
        vals = [f['nr'], TYP[f['cls']], f['message'], f['key'], f['sample'][:600],
                f['count'], f'=SUM($F$2:F{r})/SUM($F$2:$F${n + 1})', len(f['variants']),
                iso(f['first']), iso(f['last']),
                ', '.join(sorted(f['procs'])), ', '.join(sorted(f['folders'])), f['current'],
                prefill, None]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
        style_row(ws, r, len(ERR_COLS), input_cols_idx=(14, 15))
        ws.cell(row=r, column=7).number_format = '0.0%'
        ws.cell(row=r, column=6).number_format = '#,##0'
        ws.cell(row=r, column=9).number_format = 'DD.MM.YYYY'
        ws.cell(row=r, column=10).number_format = 'DD.MM.YYYY'
        ws.cell(row=r, column=4).font = grey
        ws.cell(row=r, column=5).font = grey
        dv_cat.add(ws.cell(row=r, column=14))
        ws.row_dimensions[r].height = 45
    ws.auto_filter.ref = f'A1:{get_column_letter(len(ERR_COLS))}{n + 1}'
    return ws

error_sheet('Fehler', main)
# a second DataValidation object for the second sheet (openpyxl binds one DV per sheet)
dv_cat2 = DataValidation(type='list', formula1='"' + ','.join(CATS) + '"', allow_blank=True)
dv_cat = dv_cat2
error_sheet('Business Exceptions', biz, prefill='Nicht als Fehler anzeigen')

# ── Sheet: Prozesse & Warteschlangen ────────────────────────────────────────
ws = wb.create_sheet('Prozesse & Warteschlangen')
PQ_COLS = [
    ('Typ', 14), ('Name', 48), ('Ordner', 36), ('Beschreibung', 40), ('Version', 10),
    ('Ausführungen gesamt (Jobs / Items)', 14), ('Erfolgreich (Jobs) / übrige (Items)', 14), ('Abgebrochen / System Exception', 14),
    ('Business Exceptions', 12), ('Gestoppt', 10), ('Letzte Aktivität', 14),
    ('Einbeziehen', 14), ('Kommentar', 40),
]
header(ws, PQ_COLS, input_cols=('Einbeziehen', 'Kommentar'))
dv_yn = DataValidation(type='list', formula1='"Ja,Nein"', allow_blank=True)
ws.add_data_validation(dv_yn)
rows = []
for p in sorted(d['processes'], key=lambda p: (p['folder'], p['name'])):
    rows.append(['Prozess', p['name'], p['folder'], p['description'] or None, p['version'] or None,
                 p['total'], p['successful'], p['faulted'], None, p['stopped'], iso(p['last'])])
for q in sorted(d['queues'], key=lambda q: (q['folder'], q['name'])):
    rows.append(['Warteschlange', q['name'], q['folder'], q['description'] or None, None,
                 q['total'], (q['total'] - q['failed']) if q['total'] is not None else None,
                 q['app'], q['business'], None, iso(q['last'])])
for i, vals in enumerate(rows):
    r = i + 2
    for c, v in enumerate(vals, 1):
        ws.cell(row=r, column=c, value=v)
    style_row(ws, r, len(PQ_COLS), input_cols_idx=(12, 13))
    for c in (6, 7, 8, 9, 10):
        ws.cell(row=r, column=c).number_format = '#,##0'
    ws.cell(row=r, column=11).number_format = 'DD.MM.YYYY'
    dv_yn.add(ws.cell(row=r, column=12))
ws.auto_filter.ref = f'A1:{get_column_letter(len(PQ_COLS))}{len(rows) + 1}'

# ── Sheet: Varianten (Detail) ───────────────────────────────────────────────
ws = wb.create_sheet('Varianten (Detail)')
V_COLS = [('Familie Nr', 10), ('Blatt', 20), ('Typ', 22), ('Normalisierte Meldung', 90), ('Anzahl', 9),
          ('Letztes Auftreten', 14), ('Prozesse / Warteschlangen', 40), ('Aktuelle Einstufung', 24)]
header(ws, V_COLS, freeze='A2')
r = 2
for f in fams:
    for v in f['variants']:
        vals = [f['nr'], 'Fehler' if f['cls'] != 'biz' else 'Business Exceptions', TYP[f['cls']], v['message'],
                v['count'], iso(v['lastSeen']), ', '.join(v['processes']), CUR_MAP[v['currentDashboard']]]
        for c, val in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=val)
        style_row(ws, r, len(V_COLS))
        ws.cell(row=r, column=6).number_format = 'DD.MM.YYYY'
        r += 1
ws.auto_filter.ref = f'A1:{get_column_letter(len(V_COLS))}{r - 1}'

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f'Fehler: {len(main)} families, Business: {len(biz)} families, PQ rows: {len(rows)}, variants: {r - 2}')
print('written', OUT)
